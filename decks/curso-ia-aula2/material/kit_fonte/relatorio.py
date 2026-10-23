"""Gera o relatório de conformidade com os Temas 6 e 1234 do STF em Word, a partir de um JSON.
Uso: python3 relatorio.py caso.json logo.png saida.docx   (precisa de python-docx: pip install python-docx)"""
import json, sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

VERDE, COBRE, CINZA = "064A45", "B87548", "666666"
SITUACAO = {   # rótulo, cor de fundo
    "atendido": ("Atendido", "2E7D5B"),
    "nao_comprovado": ("Alegado, não comprovado", "C98A0B"),
    "nao_encontrado": ("Não encontrei nos autos", "B3402F"),
    "dispensado": ("Dispensado", "5F7390"),
}
NOME = "Dias & Moura"


def rgb(h):
    return RGBColor.from_string(h)


def run(par, texto, tam=11, negrito=False, cor=None, italico=False):
    r = par.add_run(texto)
    r.font.name = "Georgia"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "Georgia")
    r.font.size = Pt(tam); r.bold = negrito; r.italic = italico
    if cor: r.font.color.rgb = rgb(cor)
    return r


def fundo(cel, cor):
    pr = cel._tc.get_or_add_tcPr()
    for velho in pr.findall(qn("w:shd")):
        pr.remove(velho)
    s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), cor)
    pr.append(s)


def bordas(cel, **lados):
    """lados: top/bottom/left/right = (cor, tamanho em oitavos de ponto); lado omitido fica sem borda."""
    pr = cel._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders")
    for lado in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{lado}"); v = lados.get(lado)
        e.set(qn("w:val"), "single" if v else "nil")
        if v: e.set(qn("w:sz"), str(v[1])); e.set(qn("w:color"), v[0])
        b.append(e)
    pr.append(b)


def margens(cel, cima=80, lados=140):
    pr = cel._tc.get_or_add_tcPr(); m = OxmlElement("w:tcMar")
    for k, v in (("top", cima), ("start", lados), ("bottom", cima), ("end", lados)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    pr.append(m)


def filete(par, cor=COBRE, tam=12):
    pr = par._p.get_or_add_pPr(); b = OxmlElement("w:pBdr"); e = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", str(tam)), ("space", "4"), ("color", cor)):
        e.set(qn(f"w:{k}"), v)
    b.append(e); pr.append(b)


def larguras(tab, cms):
    tab.autofit = False
    pr = tab._tbl.tblPr
    for velho in pr.findall(qn("w:tblW")) + pr.findall(qn("w:tblLayout")):
        pr.remove(velho)
    w = OxmlElement("w:tblW"); w.set(qn("w:w"), str(int(sum(cms) * 567))); w.set(qn("w:type"), "dxa"); pr.append(w)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); pr.append(lay)
    for lin in tab.rows:
        for c, w in zip(lin.cells, cms):
            c.width = Cm(w)


def campo(par, instr):
    for t, txt in (("begin", None), (None, instr), ("end", None)):
        r = run(par, "", 8.5, cor=CINZA)
        if t:
            f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), t); r._r.append(f)
        else:
            i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = txt; r._r.append(i)


def titulo_secao(doc, texto):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(16); p.paragraph_format.space_after = Pt(6)
    run(p, texto, 12.5, True, VERDE); filete(p)


def main(caso_json, logo, saida):
    c = json.load(open(caso_json, encoding="utf-8"))
    d = Document(); sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.25); sec.top_margin, sec.bottom_margin = Cm(3.0), Cm(2.2)
    d.styles["Normal"].font.name = "Georgia"; d.styles["Normal"].font.size = Pt(11)
    L = sec.page_width - sec.left_margin - sec.right_margin
    h = sec.header.paragraphs[0]; h.paragraph_format.tab_stops.add_tab_stop(L, WD_TAB_ALIGNMENT.RIGHT)
    h.add_run().add_picture(logo, height=Cm(1.4)); run(h, "\t\tDocumento de trabalho", 9.5, cor=CINZA); filete(h)
    f = sec.footer.paragraphs[0]; f.paragraph_format.tab_stops.add_tab_stop(L, WD_TAB_ALIGNMENT.RIGHT)
    run(f, f"{NOME} Advocacia · documento fictício\t\tPágina ", 8.5, cor=CINZA); campo(f, "PAGE")

    # faixa de título
    t = d.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER; larguras(t, [16.5])
    cel = t.rows[0].cells[0]; fundo(cel, VERDE); margens(cel, 220, 280); bordas(cel, bottom=(COBRE, 36))
    run(cel.paragraphs[0], "RELATÓRIO DE CONFORMIDADE", 9.5, True, "E8C9A8")
    run(cel.add_paragraph(), "Temas 6 e 1234 do STF", 22, True, "FFFFFF")
    run(cel.add_paragraph(), f"Fornecimento de {c['medicamento']}", 11.5, False, "FFFFFF", italico=True)

    # ficha do caso
    d.add_paragraph().paragraph_format.space_after = Pt(0)
    ficha = d.add_table(rows=2, cols=3); larguras(ficha, [5.5, 5.5, 5.5])
    for i, (k, v) in enumerate([("Processo", c["processo"]), ("Autor representado por", c["representacao"]),
                                ("Data da análise", c["data_analise"])]):
        a, b = ficha.rows[0].cells[i], ficha.rows[1].cells[i]
        run(a.paragraphs[0], k.upper(), 8, True, COBRE); run(b.paragraphs[0], v, 11, True, VERDE)
        for x in (a, b): margens(x, 30, 60); bordas(x)

    # veredito
    ok = sum(1 for r in c["requisitos"] if r["situacao"] in ("atendido", "dispensado"))
    total = len(c["requisitos"]); cumpriu = ok == total
    d.add_paragraph().paragraph_format.space_after = Pt(0)
    v = d.add_table(rows=1, cols=2); larguras(v, [4.2, 12.3])
    a, b = v.rows[0].cells; cor = "2E7D5B" if cumpriu else "B3402F"
    fundo(a, cor); a.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; margens(a, 140, 120)
    pa = a.paragraphs[0]; pa.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(pa, f"{ok}/{total}", 30, True, "FFFFFF")
    pa2 = a.add_paragraph(); pa2.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(pa2, "requisitos", 9.5, False, "FFFFFF")
    fundo(b, "F6EFE8"); margens(b, 140, 220); bordas(b, left=(cor, 36)); b.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    run(b.paragraphs[0], "O autor cumpriu os requisitos." if cumpriu else "O autor não cumpriu todos os requisitos.", 14, True, VERDE)
    run(b.add_paragraph(), c["resumo"], 10.5, False, "333333")

    # semáforo
    titulo_secao(d, "Tema 6: os seis requisitos")
    s = d.add_table(rows=2, cols=total); larguras(s, [16.5 / total] * total); s.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, r in enumerate(c["requisitos"]):
        rot, cor = SITUACAO[r["situacao"]]
        a, b = s.rows[0].cells[i], s.rows[1].cells[i]
        fundo(a, cor); bordas(a, left=("FFFFFF", 24), right=("FFFFFF", 24)); margens(a, 90, 40)
        pa = a.paragraphs[0]; pa.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(pa, r["id"], 20, True, "FFFFFF")
        bordas(b); margens(b, 40, 30)
        pb = b.paragraphs[0]; pb.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(pb, rot, 8, False, "444444")

    # matriz detalhada
    d.add_paragraph().paragraph_format.space_after = Pt(0)
    m = d.add_table(rows=1, cols=4)
    for cel, txt in zip(m.rows[0].cells, ("Requisito", "Situação", "Onde nos autos", "Observação")):
        fundo(cel, VERDE); margens(cel, 70, 100); run(cel.paragraphs[0], txt, 9.5, True, "FFFFFF")
    for i, r in enumerate(c["requisitos"]):
        lin = m.add_row().cells; rot, cor = SITUACAO[r["situacao"]]
        for cel in lin: margens(cel, 70, 100); bordas(cel, bottom=("D9D9D9", 4)); fundo(cel, "F4F4F4" if i % 2 else "FFFFFF")
        run(lin[0].paragraphs[0], f"({r['id']}) {r['nome']}", 9.5, True, VERDE)
        fundo(lin[1], cor); pl = lin[1].paragraphs[0]; pl.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(pl, rot, 8.5, True, "FFFFFF")
        lin[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        run(lin[2].paragraphs[0], r["onde"], 9.5); run(lin[3].paragraphs[0], r["obs"], 9.5)
    larguras(m, [4.5, 2.8, 3.7, 5.5])

    # competência
    titulo_secao(d, "Tema 1234: competência e custeio")
    cp = c["competencia"]; k = d.add_table(rows=1, cols=3); larguras(k, [5.2, 5.2, 6.1])
    for i, (rot, val) in enumerate([("Valor anual do tratamento", cp["valor_anual"]), ("210 salários mínimos", cp["limite"]),
                                    ("Onde tramita", cp["resultado"])]):
        cel = k.rows[0].cells[i]; fundo(cel, "F6EFE8" if i < 2 else VERDE); margens(cel, 140, 160)
        bordas(cel, left=("FFFFFF", 36), right=("FFFFFF", 36), top=(COBRE, 24))
        run(cel.paragraphs[0], rot.upper(), 8, True, "E8C9A8" if i == 2 else COBRE)
        run(cel.add_paragraph(), val, 14, True, "FFFFFF" if i == 2 else VERDE)
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(6); run(p, cp["nota"], 9.5, italico=True, cor=CINZA)

    # conferência humana
    titulo_secao(d, "Pontos para conferência humana")
    cf = d.add_table(rows=1, cols=1); larguras(cf, [16.5]); cel = cf.rows[0].cells[0]
    fundo(cel, "FBF3E4"); margens(cel, 140, 220); bordas(cel, left=("C98A0B", 36))
    run(cel.paragraphs[0], "Antes de usar este relatório", 10.5, True, "8A5A00")
    for item in c["conferencia"]:
        run(cel.add_paragraph(), "▪  " + item, 10)
    d.save(saida)


if __name__ == "__main__":
    main(*sys.argv[1:4])
