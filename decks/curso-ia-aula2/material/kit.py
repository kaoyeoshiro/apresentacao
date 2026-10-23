"""Monta o kit da Aula 2 em decks/curso-ia-aula2/kit/ (vai para o repositório público do curso).
Usa python-docx (uv pip install --python ~/.venvs/apresentacao/bin/python python-docx). Rode antes escritorio.py e gerar.py."""
import pathlib, shutil, subprocess, sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = pathlib.Path(__file__).parent
KIT = AQUI.parent / "kit"
ESC = AQUI / "escritorio"
VERDE, COBRE = RGBColor(0x06, 0x4A, 0x45), "B87548"
NOME = "Dias & Moura"


def fonte(run, tam=None, negrito=None, cor=None):
    run.font.name = "Georgia"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Georgia")
    if tam: run.font.size = Pt(tam)
    if negrito is not None: run.bold = negrito
    if cor is not None: run.font.color.rgb = cor


def borda(par, lado, cor, tam=12):
    pr = par._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    e = OxmlElement(f"w:{lado}")
    for k, v in (("val", "single"), ("sz", str(tam)), ("space", "4"), ("color", cor)):
        e.set(qn(f"w:{k}"), v)
    b.append(e); pr.append(b)


def campo(par, instr):
    for t, txt in (("begin", None), (None, instr), ("end", None)):
        r = par.add_run(); fonte(r, 8.5)
        if t:
            f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), t); r._r.append(f)
        else:
            i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = txt; r._r.append(i)


def sombra(celula, cor):
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), cor)
    celula._tc.get_or_add_tcPr().append(s)


def modelo(destino):
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(3.2), Cm(2.4)
    est = d.styles["Normal"]; est.font.name = "Georgia"; est.font.size = Pt(11)
    est.paragraph_format.line_spacing = 1.3
    largura = sec.page_width - sec.left_margin - sec.right_margin
    # cabeçalho: logo à esquerda, "Documento de trabalho" à direita, filete cobre
    h = sec.header.paragraphs[0]
    h.paragraph_format.tab_stops.add_tab_stop(largura, WD_TAB_ALIGNMENT.RIGHT)
    h.add_run().add_picture(str(ESC / "logo-escritorio.png"), height=Cm(1.5))
    r = h.add_run("\tDocumento de trabalho"); fonte(r, 9.5, cor=RGBColor(0x66, 0x66, 0x66))
    borda(h, "bottom", COBRE)
    # rodapé: nome e número da página
    f = sec.footer.paragraphs[0]
    f.paragraph_format.tab_stops.add_tab_stop(largura, WD_TAB_ALIGNMENT.RIGHT)
    r = f.add_run(f"{NOME} Advocacia · documento fictício\tPágina "); fonte(r, 8.5, cor=RGBColor(0x66, 0x66, 0x66))
    campo(f, "PAGE")
    # corpo
    t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_before, t.paragraph_format.space_after = Pt(14), Pt(16)
    fonte(t.add_run("{{TITULO}}"), 14, True, VERDE)
    h2 = d.add_paragraph(); h2.paragraph_format.space_before = Pt(10)
    fonte(h2.add_run("Datas e fontes"), 12, True, VERDE); borda(h2, "bottom", COBRE)
    tab = d.add_table(rows=2, cols=3)
    tab.autofit = False
    for c, (cab, larg) in enumerate((("Data", 3.2), ("Fato", 8.0), ("Fonte", 5.4))):
        for lin in tab.rows:
            lin.cells[c].width = Cm(larg)
        cel = tab.rows[0].cells[c]; sombra(cel, "064A45")
        fonte(cel.paragraphs[0].add_run(cab), 10, True, RGBColor(255, 255, 255))
        fonte(tab.rows[1].cells[c].paragraphs[0].add_run("{{c}}"), 10)
    h3 = d.add_paragraph(); h3.paragraph_format.space_before = Pt(18)
    fonte(h3.add_run("Pontos para conferência humana"), 12, True, VERDE); borda(h3, "bottom", COBRE)
    fonte(d.add_paragraph().add_run("{{CONFERENCIA}}"), 11)
    d.save(destino)


def main():
    if KIT.exists():
        shutil.rmtree(KIT)
    (KIT / "caso").mkdir(parents=True); (KIT / "escritorio").mkdir()
    gab = KIT / "gabarito"; skill = gab / "skills" / "parecer-no-padrao"; skill.mkdir(parents=True)
    for n in ("contrato", "aditivo", "notificacao"):
        shutil.copy(AQUI / "caso-ficticio" / f"{n}.pdf", KIT / "caso")
    shutil.copy(ESC / "padrao-do-escritorio.md", KIT / "escritorio")
    shutil.copy(ESC / "logo-escritorio.png", KIT / "escritorio")
    modelo(KIT / "escritorio" / "modelo.docx")
    shutil.copy(ESC / "cronologia-exemplo.md", gab / "cronologia.md")
    shutil.copy(AQUI / "kit_fonte" / "preencher.py", skill)
    (skill / "SKILL.md").write_text(SKILL, encoding="utf-8")
    (KIT / "README.md").write_text(README, encoding="utf-8")
    subprocess.run([sys.executable, str(skill / "preencher.py"), str(gab / "cronologia.md"),
                    str(KIT / "escritorio" / "modelo.docx"), str(gab / "cronologia.docx")], check=True)
    print("kit ok:", KIT)


SKILL = """---
name: parecer-no-padrao
description: Gera o cronologia.docx no padrão do escritório Dias & Moura a partir do cronologia.md da pasta. Use quando o usuário pedir a cronologia ou o parecer em Word no padrão do escritório.
---

# Parecer no padrão do escritório

1. Leia o `cronologia.md` e o `padrao-do-escritorio.md` da pasta do caso.
2. Rode `python3 preencher.py cronologia.md escritorio/modelo.docx cronologia.docx`, com o `preencher.py` desta skill.
   O script precisa de `python-docx` (`pip install python-docx`). O modelo já traz logo, cabeçalho, tabela e rodapé.
3. Confira se a tabela tem uma linha por fato, com data no formato dd/mm/aaaa, e se a seção "Pontos para conferência
   humana" lista tudo o que ficou sem fonte.
4. Salve o `cronologia.docx` na pasta e diga ao usuário o que ele precisa conferir.

Regra fixa: se um dado não está nos documentos, escreva "Não encontrei nos documentos". Nunca complete por conta própria.
"""

README = """# Kit da Aula 2: do chat à pasta

Caso fictício e escritório fictício (Dias & Moura). Nenhum dado é real.

## O que tem aqui

- `caso/`: contrato, aditivo e notificação (PDF).
- `escritorio/`: padrão de documentos, logo e `modelo.docx` já no padrão.
- `gabarito/`: a cronologia pronta, o Word final e a skill pronta. Só abra depois de tentar a sua.

## Como usar na aula

1. Baixe o kit (botão verde Code > Download ZIP) e abra a pasta `kit` no Claude, no ChatGPT Work ou no Codex.
2. Prática 1: peça a cronologia do caso, com os seis elementos do pedido e a regra "se não souber, diga".
   Salve como `cronologia.md` na pasta (copie o seu para a raiz do kit).
3. Prática 2: peça ao agente que crie uma skill chamada `parecer-no-padrao`. Ela lê o `cronologia.md`, preenche o
   `escritorio/modelo.docx` e salva o `cronologia.docx`. Rode a skill e confira o Word.
4. Travou? Compare com `gabarito/skills/parecer-no-padrao`, copie para a pasta de skills da sua ferramenta e rode.

O `modelo.docx` é preenchido por um script curto (`preencher.py`). Se o seu app gerar o Word direto, tudo bem:
o importante é o resultado seguir `escritorio/padrao-do-escritorio.md`.
"""

if __name__ == "__main__":
    main()
